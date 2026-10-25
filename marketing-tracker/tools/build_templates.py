"""Build the two Google Sheets import templates.

Usage (from the marketing-tracker folder):
    python3 tools/build_templates.py

Writes tracker_template.xlsx and seasonal_moments_template.xlsx next to index.html.
Requires openpyxl. The allowed values below must match ALLOWED in index.html.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT_DIR = Path(__file__).resolve().parent.parent

COLUMNS = [
    "id", "name", "tier", "parent_id", "vertical", "campaign_type",
    "start_date", "end_date", "status", "owner", "channels",
    "brief", "target_audience", "goal", "kpis", "key_message", "top_highlights",
    "budget_qar", "notes",
]

WIDTHS = {
    "id": 14, "name": 34, "tier": 18, "parent_id": 14, "vertical": 16,
    "campaign_type": 16, "start_date": 12, "end_date": 12, "status": 11,
    "owner": 16, "channels": 24, "brief": 48, "target_audience": 30,
    "goal": 28, "kpis": 28, "key_message": 34, "top_highlights": 40,
    "budget_qar": 12, "notes": 34,
}

TIERS = ["Seasonal Moment", "Brand Campaign", "Tactical Campaign"]
VERTICALS = [
    ("S Market", "SMK"), ("Restaurants", "RST"), ("Snoomart", "SNM"),
    ("Grocery", "GRC"), ("SLaundry", "SLD"), ("Services", "SRV"),
    ("New Ventures", "NVT"), ("Cross-vertical", "CVT"),
]
CAMPAIGN_TYPES = [
    "Hero", "Promo", "Tactical", "Occasion", "Recurrent", "USP",
    "Segmented", "Engagement", "Adoption", "Habit Builder",
]
STATUSES = ["Idea", "Planning", "Live", "Complete", "Paused"]

DATA_ROWS = 500  # validation and formatting applied to rows 2..DATA_ROWS+1

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LOCKED_FILL = PatternFill("solid", fgColor="F3F4F6")
WRAP_COLS = {"brief", "target_audience", "goal", "kpis", "key_message", "top_highlights", "notes"}


def col_letter(name):
    from openpyxl.utils import get_column_letter
    return get_column_letter(COLUMNS.index(name) + 1)


def build(path, tiers, verticals, locked=None):
    """locked: dict of column -> value to pre-fill on every data row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Initiatives"
    lists = wb.create_sheet("Lists")

    # Initiatives header
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22
    for name in COLUMNS:
        ws.column_dimensions[col_letter(name)].width = WIDTHS[name]

    last = DATA_ROWS + 1
    # Dates must export as YYYY-MM-DD. Google Sheets turns typed dates into
    # date values and publishes them in the cell's display format, so the
    # format has to be set on the column.
    for r in range(2, last + 1):
        for name in ("start_date", "end_date"):
            ws[f"{col_letter(name)}{r}"].number_format = "yyyy-mm-dd"
        ws[f"{col_letter('budget_qar')}{r}"].number_format = "0"
        for name in WRAP_COLS:
            ws[f"{col_letter(name)}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        for name, value in (locked or {}).items():
            c = ws[f"{col_letter(name)}{r}"]
            c.value = value
            c.fill = LOCKED_FILL

    # Lists tab: one column per list, plus a vertical code reference
    list_cols = [
        ("tier", tiers),
        ("vertical", [v for v, _ in verticals]),
        ("campaign_type", CAMPAIGN_TYPES),
        ("status", STATUSES),
    ]
    for i, (title, values) in enumerate(list_cols, start=1):
        lists.cell(row=1, column=i, value=title)
        for j, v in enumerate(values, start=2):
            lists.cell(row=j, column=i, value=v)
    code_col = len(list_cols) + 2
    lists.cell(row=1, column=code_col, value="vertical")
    lists.cell(row=1, column=code_col + 1, value="code")
    for j, (v, code) in enumerate(VERTICALS, start=2):
        lists.cell(row=j, column=code_col, value=v)
        lists.cell(row=j, column=code_col + 1, value=code)
    for cell in lists[1]:
        if cell.value:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
    lists.freeze_panes = "A2"
    for letter, w in zip("ABCDEFG", (18, 16, 16, 12, 3, 16, 8)):
        lists.column_dimensions[letter].width = w

    # Dropdowns referencing the Lists tab
    for i, (title, values) in enumerate(list_cols, start=1):
        from openpyxl.utils import get_column_letter
        src = get_column_letter(i)
        dv = DataValidation(
            type="list",
            formula1=f"=Lists!${src}$2:${src}${len(values) + 1}",
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Not an allowed value",
            error=f"Pick a {title} from the dropdown (see the Lists tab).",
        )
        dv.add(f"{col_letter(title)}2:{col_letter(title)}{last}")
        ws.add_data_validation(dv)

    wb.save(path)
    print(f"wrote {path}")


build(OUT_DIR / "tracker_template.xlsx", TIERS[1:], VERTICALS)
build(
    OUT_DIR / "seasonal_moments_template.xlsx",
    TIERS[:1],
    VERTICALS[-1:],
    locked={"tier": "Seasonal Moment", "vertical": "Cross-vertical"},
)
